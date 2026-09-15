"""Text extraction for PDF and DOCX files.

Dispatches on the file suffix. Raises ExtractionError for
unsupported types, encrypted PDFs, or unreadable files.
"""

from pathlib import Path


class ExtractionError(Exception):
    """Raised when text cannot be extracted from a document."""


def extract_text(path: Path) -> str:
    """Extract raw text from a PDF or DOCX file at *path*."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    raise ExtractionError(f"Unsupported file type: {ext or '(none)'}")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not read PDF file: {exc}") from exc
    if getattr(reader, "is_encrypted", False):
        try:
            if reader.decrypt("") == 0:
                raise ExtractionError("Encrypted PDF is not supported.")
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError(f"Could not read encrypted PDF: {exc}") from exc
    try:
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:
        raise ExtractionError(f"Could not extract PDF text: {exc}") from exc
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX file: {exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
