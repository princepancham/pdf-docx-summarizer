"""Extraction tests using synthetic fixtures (no network, no real docs)."""

from pathlib import Path

import pytest

from backend.extract import ExtractionError, extract_text
from tests.pdf_fixture import minimal_pdf_bytes


def test_extract_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    from docx import Document

    src = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Hello DOCX world")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell one"
    table.rows[0].cells[1].text = "cell two"
    doc.save(str(src))

    text = extract_text(src)
    assert "Hello DOCX world" in text
    assert "cell one" in text
    assert "cell two" in text


def test_extract_pdf(tmp_path: Path) -> None:
    src = tmp_path / "sample.pdf"
    src.write_bytes(minimal_pdf_bytes())
    assert "Hello PDF world" in extract_text(src)


def test_extract_unsupported_type(tmp_path: Path) -> None:
    src = tmp_path / "notes.txt"
    src.write_text("hello", encoding="utf-8")
    with pytest.raises(ExtractionError):
        extract_text(src)


def test_extract_missing_file(tmp_path: Path) -> None:
    with pytest.raises(ExtractionError):
        extract_text(tmp_path / "nope.pdf")
