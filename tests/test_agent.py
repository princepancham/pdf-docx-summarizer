"""Agent tests: profile, plan/fallback, execute, evaluate, repair (mocked LLM)."""

import dataclasses
import logging

import pytest

import backend.agent as agent_module
import backend.llm as llm_module
from backend.agent import (
    deterministic_plan,
    evaluate,
    plan_strategy,
    profile_text,
    run,
)
from backend.config import settings as prod_settings, validate_settings
from backend.llm import LLMError


@pytest.fixture()
def iso(monkeypatch: pytest.MonkeyPatch):
    """Isolated agent/llm settings with default chunk budgets."""
    isolated = dataclasses.replace(prod_settings, openrouter_api_key="test-key")
    monkeypatch.setattr(agent_module, "settings", isolated)
    monkeypatch.setattr(llm_module, "settings", isolated)
    return isolated


def _plan_unavailable(monkeypatch: pytest.MonkeyPatch) -> None:
    def boom(system: str, user: str) -> dict:
        raise LLMError("AI service returned an invalid response.", 502)

    monkeypatch.setattr(llm_module, "chat_json", boom)


# -- profile -------------------------------------------------------------


def test_profile_short_sparse(iso) -> None:
    profile = profile_text("Hello world. Short note here.")
    assert profile.density == "sparse"
    assert profile.paragraphs <= 3
    assert profile.headings == 0


def test_profile_dense_long(iso) -> None:
    text = "\n\n".join("Wordy paragraph content here. " * 30 for _ in range(10))
    profile = profile_text(text)
    assert profile.density == "dense"
    assert profile.avg_para_len > 800


def test_profile_headings_detected(iso) -> None:
    text = "# Introduction\n\nSome body text here.\n\n## Methods\n\nMore text."
    assert profile_text(text).headings >= 2


def test_profile_uses_text_signals_only(iso) -> None:
    profile = profile_text("plain text without any structure")
    assert not hasattr(profile, "table_cells")
    assert not hasattr(profile, "table_density")


# -- deterministic plan --------------------------------------------------


def test_deterministic_short_is_direct(iso) -> None:
    assert deterministic_plan(profile_text("tiny doc")).strategy == "direct"


def test_deterministic_dense_is_key_points_first(iso) -> None:
    text = "\n\n".join("Wordy paragraph content here. " * 30 for _ in range(10))
    assert deterministic_plan(profile_text(text)).strategy == "key_points_first"


def test_deterministic_long_normal_is_map_reduce(iso) -> None:
    text = "\n\n".join(f"Paragraph {i} with some normal text. " * 20 for i in range(30))
    profile = profile_text(text)
    assert profile.chars > 4000
    assert deterministic_plan(profile).strategy == "map_reduce"


# -- planning call -------------------------------------------------------


def test_plan_json_selects_strategy(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        llm_module,
        "chat_json",
        lambda system, user: {
            "strategy": "key_points_first",
            "focus": "methods",
            "max_sections": 5,
        },
    )
    plan, fallback = plan_strategy(profile_text("x" * 9000))
    assert plan.strategy == "key_points_first"
    assert plan.max_sections == 5
    assert fallback is False


def test_plan_garbage_falls_back(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _plan_unavailable(monkeypatch)
    plan, fallback = plan_strategy(profile_text("tiny"))
    assert plan.strategy == "direct"
    assert fallback is True


def test_plan_bad_strategy_falls_back(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        llm_module, "chat_json", lambda system, user: {"strategy": "nonsense"}
    )
    plan, fallback = plan_strategy(profile_text("tiny"))
    assert fallback is True
    assert plan.strategy == "direct"


def test_plan_missing_key_propagates(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    isolated = dataclasses.replace(prod_settings, openrouter_api_key=None)
    monkeypatch.setattr(llm_module, "settings", isolated)
    monkeypatch.setattr(agent_module, "settings", isolated)
    with pytest.raises(LLMError) as exc_info:
        plan_strategy(profile_text("tiny"))
    assert exc_info.value.status_code == 503


# -- execute + evaluate + repair -----------------------------------------


def _good_chat(words: str):
    def fake(messages):
        return (
            f"Summary about {words} with plenty of extra descriptive text "
            f"to satisfy the length gate comfortably. {words} matters."
        )

    return fake


def test_run_direct_single_call(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _plan_unavailable(monkeypatch)
    text = (
        "The honeybee waggle dance communicates direction. "
        "The honeybee waggle dance communicates distance. "
        "Researchers study honeybee behavior carefully."
    )
    monkeypatch.setattr(
        llm_module,
        "_chat",
        lambda messages: (
            "Honeybee waggle dance communicates direction and distance. "
            "Researchers study honeybee behavior carefully in detail."
        ),
    )
    result = run(text)
    assert result.strategy == "direct"
    assert result.calls == 2  # plan attempt + execute
    assert 0.0 <= result.quality_score <= 1.0
    assert result.chunks == 1
    assert result.truncated is False


def test_run_map_reduce_counts_calls(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _plan_unavailable(monkeypatch)
    text = "\n\n".join(f"Paragraph {i} about Trains and bridges. " * 20 for i in range(30))
    seen: list[str] = []

    def fake(messages):
        seen.append(messages[-1]["content"])
        if messages[-1]["content"].startswith("Combine"):
            return (
                "Combined summary about trains and bridges. " * 12
                + "Trains bridges paragraph section content."
            )
        return (
            "Partial about trains and bridges, long enough to be useful here."
        )

    monkeypatch.setattr(llm_module, "_chat", fake)
    result = run(text)
    assert result.strategy == "map_reduce"
    assert result.chunks > 1
    assert result.calls == 1 + result.chunks + 1  # plan + parts + combine


def test_run_key_points_first_extracts_points(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        llm_module,
        "chat_json",
        lambda system, user: {"strategy": "key_points_first", "max_sections": 12},
    )
    text = "\n\n".join(f"Section {i} about solar panels. " * 40 for i in range(8))

    def fake(messages):
        content = messages[-1]["content"]
        if content.startswith("List the key points"):
            return "- solar panels convert sunlight\n- installation requires roof space\n- maintenance is minimal upkeep"
        return (
            "Combined overview of solar panels covering sunlight conversion "
            "and roof installation and minimal upkeep in good detail."
        )

    monkeypatch.setattr(llm_module, "_chat", fake)
    result = run(text)
    assert result.strategy == "key_points_first"
    assert len(result.key_points) >= 3
    assert any("solar" in p for p in result.key_points)


def test_repair_retry_bounded(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _plan_unavailable(monkeypatch)
    calls: list[str] = []
    monkeypatch.setattr(
        llm_module, "_chat", lambda messages: calls.append("x") or "x"
    )
    result = run("Some reasonable document text about gardening tools.")
    # plan + initial execute + exactly 1 repair execute
    assert result.calls == 3
    assert any("repair" in n for n in result.notes)
    assert 0.0 <= result.quality_score < 1.0


def test_repair_disabled_by_config(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _plan_unavailable(monkeypatch)
    isolated = dataclasses.replace(prod_settings, agent_max_repairs=0)
    monkeypatch.setattr(agent_module, "settings", isolated)
    monkeypatch.setattr(llm_module, "_chat", lambda messages: "x")
    result = run("Some reasonable document text about gardening tools.")
    assert result.calls == 2  # plan + execute, no repair


def test_execute_failure_carries_strategy(
    iso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _plan_unavailable(monkeypatch)

    def boom(messages):
        raise LLMError("AI service timed out.", 504)

    monkeypatch.setattr(llm_module, "_chat", boom)
    with pytest.raises(LLMError) as exc_info:
        run("A short but valid document text here.")
    assert exc_info.value.status_code == 504
    assert getattr(exc_info.value, "strategy", None) == "direct"


def test_evaluate_gates() -> None:
    source = " ".join(f"uniqueword{i} common text filler" for i in range(20))
    summary = (
        "A proper summary mentioning uniqueword1 uniqueword2 uniqueword3 "
        "uniqueword4 uniqueword5 uniqueword6 with sufficient length padding "
        "to satisfy the relative length gate for this source document."
    )
    score, notes = evaluate(
        summary, [], source, strategy="direct", truncated=False
    )
    assert score == 1.0
    assert notes == []

    score, notes = evaluate("tiny", [], source, strategy="direct", truncated=False)
    assert score < 1.0
    assert notes


# -- polish: logging redaction, config validation ------------------------


def test_summarize_logs_redacted(
    tmp_path, monkeypatch: pytest.MonkeyPatch, caplog
) -> None:
    """Full mocked summarize via HTTP: logs exist but never carry the key."""
    import io
    import logging

    from docx import Document
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    import backend.agent as agent_module
    import backend.main as main_module
    from backend.database import Base, get_db
    from backend.main import app

    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    isolated = dataclasses.replace(
        prod_settings, upload_dir=upload_dir, openrouter_api_key="test-key"
    )
    monkeypatch.setattr(main_module, "settings", isolated)
    monkeypatch.setattr(llm_module, "settings", isolated)
    monkeypatch.setattr(agent_module, "settings", isolated)
    monkeypatch.setattr(
        llm_module, "chat_json", lambda system, user: {"strategy": "direct"}
    )
    monkeypatch.setattr(llm_module, "_chat", _good_chat("honeybee"))

    engine = create_engine(
        f"sqlite:///{tmp_path / 'test.db'}",
        connect_args={"check_same_thread": False},
    )
    TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    from backend import models  # noqa: F401

    Base.metadata.create_all(bind=engine)

    def override_get_db():
        session = TestingSession()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db
    try:
        doc = Document()
        doc.add_paragraph("The honeybee waggle dance communicates direction.")
        buf = io.BytesIO()
        doc.save(buf)
        with caplog.at_level(logging.INFO, logger="pdf_summarizer"):
            response = TestClient(app).post(
                "/api/documents/summarize",
                files={
                    "file": (
                        "bees.docx",
                        buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document",
                    )
                },
            )
        assert response.status_code == 200, response.text
        assert caplog.records, "expected summarize log records"
        for record in caplog.records:
            assert "test-key" not in record.getMessage()
            assert "Bearer" not in record.getMessage()
    finally:
        app.dependency_overrides.clear()


def test_validate_settings_rejects_bad_values() -> None:
    good = dataclasses.replace(prod_settings)
    validate_settings(good)
    bad = dataclasses.replace(prod_settings, chunk_overlap=4000, chunk_chars=4000)
    try:
        validate_settings(bad)
    except RuntimeError as exc:
        assert "CHUNK_OVERLAP" in str(exc)
    else:
        raise AssertionError("expected RuntimeError")
    bad2 = dataclasses.replace(prod_settings, agent_min_coverage=0.0)
    try:
        validate_settings(bad2)
    except RuntimeError as exc:
        assert "AGENT_MIN_COVERAGE" in str(exc)
    else:
        raise AssertionError("expected RuntimeError")
